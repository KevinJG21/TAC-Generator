from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE


DOWNLOADS = Path(r"C:\Users\yash\Downloads")
SOURCE_PPT = DOWNLOADS / "24BCI0123_Yash Singh Jhinkwan_Review 2.pptx"
OUTPUT_PPT = DOWNLOADS / "24BCI0123_Yash Singh Jhinkwan_Review 2 (with CFG).pptx"
CFG_IMAGE = Path(r"C:\Users\yash\Desktop\COMPILER  DESIGN\Screenshot 2026-10-16 130048.png")

NAVY = RGBColor(0x17, 0x32, 0x4D)
TEAL = RGBColor(0x00, 0xA6, 0xA6)
ORANGE = RGBColor(0xF2, 0x8C, 0x28)
LIGHT = RGBColor(0xF3, 0xF7, 0xFA)
MID = RGBColor(0xD9, 0xE5, 0xEC)
INK = RGBColor(0x1D, 0x29, 0x39)
MUTED = RGBColor(0x52, 0x60, 0x6D)


def add_text(slide, text, x, y, w, h, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.05)
    frame.margin_right = Inches(0.05)
    para = frame.paragraphs[0]
    para.alignment = align
    run = para.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def add_rect(slide, x, y, w, h, fill, line=None, rounded=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def add_bullet(slide, text, y):
    add_rect(slide, 0.95, y + 0.09, 0.12, 0.12, TEAL, TEAL, True)
    add_text(slide, text, 1.22, y, 7.5, 0.48, 16, INK)


def main():
    if not SOURCE_PPT.exists():
        raise FileNotFoundError(SOURCE_PPT)
    if not CFG_IMAGE.exists():
        raise FileNotFoundError(CFG_IMAGE)

    prs = Presentation(str(SOURCE_PPT))
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    add_rect(slide, 0, 0, 13.333, 0.12, TEAL, TEAL, False)
    add_text(slide, "VISUAL EVIDENCE", 0.65, 0.35, 4.5, 0.25, 10, TEAL, True)
    add_text(slide, "CFG generated from the Review-2 TAC", 0.65, 0.68, 8.4, 0.55, 27, NAVY, True)
    add_text(slide, "06A", 12.0, 0.45, 0.7, 0.3, 14, ORANGE, True, PP_ALIGN.RIGHT)

    add_text(slide, "Why this belongs in the presentation", 0.85, 1.55, 7.6, 0.35, 18, NAVY, True)
    add_bullet(slide, "Makes the TAC control-flow structure visible instead of only showing a text table.", 2.15)
    add_bullet(slide, "Shows ENTRY → EXIT flow, conditional branches, and the loop back-edge.", 2.9)
    add_bullet(slide, "Matches the Review-2 prototype result: nine basic blocks from the generated TAC.", 3.65)
    add_bullet(slide, "Supports the novelty and implementation evidence without changing the existing wording.", 4.4)

    add_rect(slide, 0.85, 5.45, 7.85, 0.9, LIGHT, MID, True)
    add_text(slide, "Prototype evidence", 1.15, 5.68, 1.75, 0.25, 13, TEAL, True)
    add_text(slide, "9 basic blocks  •  branch edges  •  join edges  •  loop back-edge", 2.95, 5.65, 5.4, 0.3, 15, NAVY, True)

    add_rect(slide, 9.35, 1.2, 3.2, 5.95, RGBColor(0xF8, 0xFA, 0xFB), MID, True)
    slide.shapes.add_picture(str(CFG_IMAGE), Inches(9.48), Inches(1.3), height=Inches(5.75))
    add_text(slide, "TAC control-flow graph", 9.6, 6.82, 2.8, 0.2, 9, MUTED, False, PP_ALIGN.CENTER)
    add_text(slide, "BCSE307P  •  24BCI0123  •  REVIEW 2", 0.65, 7.1, 5.0, 0.2, 8, MUTED)
    add_text(slide, "75% implementation milestone", 9.0, 7.1, 3.7, 0.2, 8, MUTED, False, PP_ALIGN.RIGHT)

    # Place the visual-evidence slide immediately after the existing CFG slide.
    target_index = len(prs.slides) - 1
    for i, existing in enumerate(prs.slides):
        text = " ".join(shape.text for shape in existing.shapes if hasattr(shape, "text"))
        if "From TAC to basic blocks and CFG" in text:
            target_index = i + 1
            break
    slide_id = prs.slides._sldIdLst[-1]
    prs.slides._sldIdLst.remove(slide_id)
    prs.slides._sldIdLst.insert(target_index, slide_id)
    prs.save(str(OUTPUT_PPT))
    print(OUTPUT_PPT)


if __name__ == "__main__":
    main()
