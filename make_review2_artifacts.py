from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches as PI, Pt as PP
from pptx.dml.color import RGBColor as PRGB
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import fitz

OUT = Path(r"C:\Users\yash\Downloads")
DOCX = OUT / "24BCI0123_Yash Singh Jhinkwan_Review 2 Report.docx"
PDF = OUT / "24BCI0123_Yash Singh Jhinkwan_Review 2 Report.pdf"
PPT = OUT / "24BCI0123_Yash Singh Jhinkwan_Review 2.pptx"
NAVY, TEAL, ORANGE, LIGHT, MID, INK, MUTED = "17324D", "00A6A6", "F28C28", "F3F7FA", "D9E5EC", "1D2939", "52606D"

def shade(cell, fill):
    tc = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tc.append(shd)

def cell_text(cell, text, bold=False, color=INK, size=8.7):
    cell.text = ""; p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text)); r.bold = bold; r.font.name = "Aptos"; r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(color)

def table(doc, heads, rows):
    t = doc.add_table(rows=1, cols=len(heads)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = "Table Grid"
    for i, h in enumerate(heads): cell_text(t.rows[0].cells[i], h, True, "FFFFFF"); shade(t.rows[0].cells[i], NAVY)
    for row in rows:
        cs = t.add_row().cells
        for i, value in enumerate(row): cell_text(cs[i], value); shade(cs[i], LIGHT if len(t.rows) % 2 == 0 else "FFFFFF")
    doc.add_paragraph().paragraph_format.space_after = Pt(2); return t

def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level); p.style.font.name = "Aptos Display"; p.style.font.color.rgb = RGBColor.from_string(NAVY if level == 1 else TEAL); return p

def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(2); p.add_run(item)

def code(doc, text):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(.25); p.paragraph_format.right_indent = Inches(.25)
    r = p.add_run(text); r.font.name = "Consolas"; r.font.size = Pt(8.5); r.font.color.rgb = RGBColor.from_string("123B4A"); return p

def make_docx():
    d = Document(); sec = d.sections[0]; sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65); sec.left_margin = Inches(.75); sec.right_margin = Inches(.75)
    d.styles["Normal"].font.name = "Aptos"; d.styles["Normal"].font.size = Pt(10); d.styles["Normal"].font.color.rgb = RGBColor.from_string(INK)
    h = sec.header.paragraphs[0]; h.text = "BCSE307P  |  PROJECT REVIEW 2  |  16 OCTOBER 2026"; h.alignment = WD_ALIGN_PARAGRAPH.RIGHT; h.runs[0].font.size = Pt(8)
    f = sec.footer.paragraphs[0]; f.alignment = WD_ALIGN_PARAGRAPH.CENTER; f.add_run("24BCI0123  •  Three Address Code Generator  •  Review 2").font.size = Pt(8)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(28); r = p.add_run("THREE ADDRESS CODE GENERATOR"); r.bold = True; r.font.size = Pt(26); r.font.color.rgb = RGBColor.from_string(NAVY)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run("PROJECT REVIEW 2  |  75% IMPLEMENTATION MILESTONE"); r.bold = True; r.font.size = Pt(15); r.font.color.rgb = RGBColor.from_string(TEAL)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run("BCSE307P Compiler Design Lab\nSchool of Computer Science and Engineering\nVIT | Fall Semester 2026–27").font.size = Pt(12)
    table(d, ["Report detail", "Value"], [("Assessment", "9 — Project Review 2"), ("Faculty", "Dr. Bhawana Tyagi"), ("Team", "Yash Singh Jhinkwan (24BCI0123)\nFaaz Shabeer (24BCE2754)\nKevin Gilbert (24BCE2843)"), ("Submission date", "16 October 2026"), ("Prototype", "Python CLI: tacgen")])
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run("Milestone statement: the submission demonstrates the implemented 75% review scope; the final 25% remains explicitly planned and is not presented as complete."); r.italic = True; r.font.color.rgb = RGBColor.from_string(ORANGE)
    d.add_page_break()
    heading(d, "1. Review-2 requirement mapping")
    d.add_paragraph("Assessment IX requires a Project Review 2 submission with 75% prototype progress, novelty, a progress report, and a presentation with Q&A. This report is organized around those four evaluation points.")
    table(d, ["Rubric item", "Evidence in this submission", "Status"], [("75% implementation / prototype — 5M", "Lexer, parser/AST, semantic checks, TAC, Boolean short-circuiting, control flow, optimizer, quads/triples, VM, and Review-2 CFG analysis.", "Demonstrated"), ("Novelty — 2M", "Short-circuit Boolean TAC plus side-by-side TAC, quadruple, triple, VM, and CFG views for one small language.", "Demonstrated"), ("Documentation / progress report — 1M", "Architecture, implementation boundary, evidence, tests, limitations, and next steps.", "Included"), ("Presentation & Q&A — 2M", "Companion PPT covers the problem, design, demo path, results, novelty, and viva prompts.", "Prepared")])
    heading(d, "2. Executive summary")
    d.add_paragraph("The project is a small compiler front end that translates a source language into Three Address Code (TAC), an intermediate representation in which each instruction has at most one operator and explicit temporary or label names. Review 2 advances the Review-1 prototype with a reproducible control-flow-graph analysis layer and an evidence-driven demonstration.")
    bullets(d, ["Implemented scope: lexical analysis, recursive-descent parsing, AST construction, semantic checks, TAC generation, short-circuit Boolean translation, if/else/while, optimization, quadruples, triples, and a TAC virtual machine.", "Review-2 addition: basic-block detection and directed CFG successor edges from the generated TAC listing.", "Verification: 23 automated tests pass, including 300 randomized programs checked in four optimizer/Boolean-mode combinations; the Review-2 demo reports AST-vs-TAC verification PASS.", "Boundary: SSA construction, functions/arrays/types, machine-code generation, and final release packaging remain outside this 75% milestone."])
    heading(d, "3. Problem, objectives, and novelty")
    d.add_paragraph("Directly translating every source-language construct to every target machine creates unnecessary coupling. TAC provides a compact intermediate form that exposes evaluation order, temporaries, labels, branches, and data flow so later compiler stages can operate on a uniform representation.")
    table(d, ["Objective", "Review-2 evidence"], [("Define the source language", "Assignments, print, arithmetic, relational and Boolean expressions, blocks, if/else, and while."), ("Build the front end", "Regex lexer with line/column errors and recursive-descent parser with precedence levels."), ("Generate multiple IR views", "TAC listing, quadruples, triples, labels, and temporary values."), ("Preserve Boolean semantics", "Jump mode uses short-circuit flow; value mode materializes 0/1 values."), ("Verify behavior", "AST interpreter and TAC VM are compared for outputs and final environments.")])
    d.add_paragraph("Novelty in the prototype is the unified learning-oriented inspection path: one source program can produce tokens, AST, TAC, quadruples, triples, symbol table, CFG, and VM verification from one CLI command. The CFG view connects textbook TAC to the control-flow structures that later enable SSA and optimization.")
    heading(d, "4. System design")
    code(d, "source text → lexer → parser / AST → semantic checks → TAC generator → optimizer → {TAC, quads, triples, CFG} → TAC VM\n                                           ↘ AST interpreter ────────────────┘  (cross-check)")
    table(d, ["Stage", "Module(s)", "Delivered behavior"], [("Lexical", "lexer.py", "Tokens, comments, reserved names, line/column diagnostics."), ("Syntax", "parser.py, ast_nodes.py", "Recursive descent with arithmetic, relational, Boolean, unary, and statement grammar."), ("Semantic + codegen", "codegen.py", "Use-before-assignment checks, symbol types, temporaries, labels, and jumps."), ("Optimization", "optimizer.py", "Constant folding, copy elimination, redundant goto and unused-label cleanup."), ("Representations", "formats.py, cfg.py", "TAC, quadruples, triples, symbol table, and CFG table."), ("Verification", "runtime.py", "AST interpreter and TAC VM execute the same program and are compared.")])
    heading(d, "5. Implementation status — 75% milestone")
    d.add_paragraph("The progress boundary is intentional: the core front end, IR generation, execution, and control-flow analysis are implemented and demonstrable. The remaining 25% is listed as future work in Section 8.")
    table(d, ["Area", "Review-2 status", "Evidence"], [("Lexer + parser + AST", "Complete for scoped language", "Token tests, precedence tests, syntax diagnostics."), ("Semantic analysis", "Implemented", "Use-before-assignment, inferred symbol types, division-by-zero checks."), ("TAC + control flow", "Implemented", "Arithmetic, relational, Boolean short-circuit, if/else, while, labels."), ("IR inspection", "Implemented", "TAC, quadruple, triple, symbol-table, and CFG output modes."), ("Verification", "Implemented for scoped language", "23 tests; randomized equivalence checks; VM output comparison."), ("Remaining 25%", "Planned, not claimed", "SSA sketch, richer language features, packaging, and final benchmark/report polish.")])
    code(d, "python -m unittest discover -s tests -v\n# 23 tests ... OK\n\npython -m tacgen examples/07_review2_demo.src --cfg --run\n# Verification against AST interpreter: PASS\n# Final variables: n=5, sum=43, i=6")
    heading(d, "6. Prototype demonstration")
    d.add_paragraph("The Review-2 example combines a while loop, a short-circuit && condition, an if/else branch, modulo, arithmetic temporaries, and output. It is small enough for a viva while exercising the major implemented paths.")
    code(d, "n = 5;\nsum = 0;\ni = 1;\nwhile (i <= n && sum < 100) {\n    if (i % 2 == 0 || i == 5) { sum = sum + i * i; }\n    else { sum = sum - 1; }\n    i = i + 1;\n}\nprint sum;")
    code(d, "L1:\n    if i > n goto L2\n    if sum >= 100 goto L2\n    t1 = i % 2\n    if t1 == 0 goto L3\n    if i != 5 goto L4\nL3:  t2 = i * i;  t3 = sum + t2;  sum = t3\nL2:  print sum\n\nCFG summary: 9 basic blocks; branch edges and loop back-edge detected.\nVM output: 43 | AST-vs-TAC: PASS")
    d.add_paragraph("In jump mode, the right side of && is reached only when the left side is true. The same principle is used for || and conditionals, producing explicit labels and conditional jumps instead of flattening every Boolean into an eager value.")
    heading(d, "7. Verification and quality checks")
    table(d, ["Check", "Result"], [("Automated unit/integration suite", "23 tests passed"), ("Randomized semantic equivalence", "300 generated programs × 4 modes passed"), ("Review-2 control-flow demo", "CFG emitted; 9 blocks; VM output 43"), ("AST interpreter vs TAC VM", "PASS on Review-2 demo and test suite"), ("Error paths", "Lexical, syntax, semantic, and runtime division-by-zero cases covered")])
    d.add_paragraph("The generated CFG is derived from the same TAC that is executed by the VM, so the analysis output is tied directly to the demonstrated intermediate representation rather than being a separate mock diagram.")
    heading(d, "8. Remaining 25% and next steps")
    bullets(d, ["SSA sketch and phi-placement discussion over the CFG.", "Richer language support: functions, arrays, explicit declarations, and more precise type rules.", "Additional optimization and benchmark reporting over larger source suites.", "Final packaging, screenshots, and release documentation for the complete project submission."])
    d.add_paragraph("These items are deliberately marked as pending. The Review-2 artifact demonstrates the requested 75% stage and does not present the project as a finished 100% implementation.")
    heading(d, "9. Viva / Q&A preparation")
    table(d, ["Likely question", "Answer cue"], [("Why use TAC?", "It makes evaluation order, temporaries, branches, and labels explicit while decoupling the front end from a target machine."), ("How does && short-circuit?", "The left condition jumps to false when false; otherwise control falls through to the right condition."), ("Quadruples vs triples?", "Quadruples name results explicitly; triples refer to earlier instruction indices."), ("What is a basic block?", "A maximal straight-line sequence with one entry; CFG edges connect possible transfers of control."), ("Why is this 75%?", "The core compiler pipeline and analysis are demonstrable; SSA, richer features, and final release work remain.")])
    heading(d, "10. References")
    bullets(d, ["A. V. Aho, M. S. Lam, R. Sethi, J. D. Ullman, Compilers: Principles, Techniques, and Tools, 2nd ed., Pearson, 2006.", "K. D. Cooper and L. Torczon, Engineering a Compiler, 2nd ed., Morgan Kaufmann, 2011.", "A. W. Appel, Modern Compiler Implementation in C, Cambridge University Press, 1998.", "C. Lattner and V. Adve, LLVM: A Compilation Framework for Lifelong Program Analysis and Transformation, CGO, 2004.", "R. Cytron et al., Efficiently Computing Static Single Assignment Form and the Control Dependence Graph, ACM TOPLAS, 1991."])
    d.save(DOCX)

def rgb(x): return PRGB.from_string(x)
def box(s, text, x, y, w, h, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, font="Aptos"):
    b = s.shapes.add_textbox(PI(x), PI(y), PI(w), PI(h)); tf = b.text_frame; tf.clear(); tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align; r = p.add_run(); r.text = text; r.font.name = font; r.font.size = PP(size); r.font.bold = bold; r.font.color.rgb = rgb(color); return b
def rect(s, x, y, w, h, fill, line=None, rounded=True):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, PI(x), PI(y), PI(w), PI(h)); sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill); sh.line.color.rgb = rgb(line or fill); return sh
def slide_title(s, kicker, heading, n):
    rect(s, 0, 0, 13.333, .12, TEAL, TEAL, False); box(s, kicker.upper(), .65, .35, 8, .25, 10, TEAL, True); box(s, heading, .65, .68, 11.5, .6, 27, NAVY, True); box(s, f"0{n}", 12.1, .45, .55, .3, 14, ORANGE, True, PP_ALIGN.RIGHT)
def foot(s): box(s, "BCSE307P  •  24BCI0123  •  REVIEW 2", .65, 7.1, 5, .2, 8, MUTED); box(s, "75% implementation milestone", 9, 7.1, 3.7, .2, 8, MUTED, False, PP_ALIGN.RIGHT)
def bullets_ppt(s, items, x, y, w, size=16, gap=.58):
    for i, item in enumerate(items): rect(s, x, y+i*gap+.08, .12, .12, TEAL, TEAL); box(s, item, x+.25, y+i*gap, w-.25, gap, size, INK)

def make_ppt():
    p = Presentation(); p.slide_width = PI(13.333); p.slide_height = PI(7.5); blank = p.slide_layouts[6]
    s=p.slides.add_slide(blank); rect(s,0,0,13.333,7.5,NAVY,NAVY,False); rect(s,0,0,13.333,.18,TEAL,TEAL,False); box(s,"BCSE307P  |  PROJECT REVIEW 2",.8,.75,8,.3,13,"8FE3DD",True); box(s,"Three Address\nCode Generator",.8,1.45,8.9,1.55,34,"FFFFFF",True); box(s,"75% implementation / prototype milestone",.85,3.35,7.5,.4,19,ORANGE,True); box(s,"Yash Singh Jhinkwan 24BCI0123\nFaaz Shabeer 24BCE2754  •  Kevin Gilbert 24BCE2843\nDr. Bhawana Tyagi  |  16 October 2026",.85,5.75,8,.8,13,"E7EEF4"); rect(s,10.1,1.55,2.1,2.1,TEAL,TEAL); box(s,"TAC",10.1,2,2.1,.6,32,"FFFFFF",True,PP_ALIGN.CENTER); box(s,"source → IR → verify",9.75,4.15,2.8,.3,13,"8FE3DD",False,PP_ALIGN.CENTER)
    s=p.slides.add_slide(blank); slide_title(s,"Review brief","What Review 2 asks us to prove",2); foot(s); box(s,"Assessment IX",.75,1.55,2.1,.3,13,TEAL,True); box(s,"75% prototype progress",.75,1.95,4,.55,25,NAVY,True); rect(s,.8,2.8,6.2,.28,MID,MID); rect(s,.8,2.8,4.65,.28,TEAL,TEAL,False); box(s,"75% implemented",5.65,2.68,1.35,.35,12,TEAL,True,PP_ALIGN.RIGHT); bullets_ppt(s,["5M  implementation / prototype","2M  novelty","1M  documentation / progress","2M  presentation & Q&A"],.85,3.5,5.7,17,.62); rect(s,7.65,1.55,4.75,4.65,LIGHT,MID); box(s,"Milestone boundary",8.05,1.95,3.8,.35,16,NAVY,True); box(s,"Implemented and demonstrated",8.05,2.55,3.8,.3,13,TEAL,True); bullets_ppt(s,["front end + AST","TAC / quads / triples","short-circuit control flow","optimizer + VM","basic blocks + CFG"],8.05,3,3.8,14,.48); box(s,"Remaining work stays visible; it is not presented as complete.",8.05,5.55,3.8,.45,12,ORANGE,True)
    s=p.slides.add_slide(blank); slide_title(s,"The problem","Why an intermediate representation matters",3); box(s,"Source program",.8,1.45,2,.3,13,MUTED,True); rect(s,.8,1.9,3.1,2.6,LIGHT,MID); box(s,"x = a + b * c;",1.1,2.35,2.5,.4,20,NAVY,True,PP_ALIGN.CENTER,"Consolas"); box(s,"precedence\nbranches\nBoolean short-circuit",1.1,3.05,2.5,.9,15,MUTED,False,PP_ALIGN.CENTER); box(s,"→",4.25,2.55,.7,.7,30,TEAL,True,PP_ALIGN.CENTER); box(s,"Three Address Code",5.1,1.45,3,.3,13,MUTED,True); rect(s,5.1,1.9,3.6,2.6,NAVY,NAVY); box(s,"t1 = b * c\nt2 = a + t1\nx = t2",5.45,2.35,2.9,1.1,19,"FFFFFF",True,PP_ALIGN.CENTER,"Consolas"); box(s,"one operator per instruction\nexplicit temporaries and labels",5.45,3.75,2.9,.55,13,"BFE9E5",False,PP_ALIGN.CENTER); box(s,"→",9.05,2.55,.7,.7,30,TEAL,True,PP_ALIGN.CENTER); box(s,"What we can inspect",9.9,1.45,2.7,.3,13,MUTED,True); bullets_ppt(s,["TAC listing","quadruples / triples","CFG + VM result"],9.95,2,2.5,16,.65); box(s,"Novelty: one command exposes the full learning path from source to verified IR.",1,5.65,11.2,.45,17,ORANGE,True,PP_ALIGN.CENTER)
    s=p.slides.add_slide(blank); slide_title(s,"Architecture","The implemented pipeline",4); foot(s); stages=[("1","Lexer","tokens"),("2","Parser","AST"),("3","Semantic","symbols"),("4","TAC","temps + labels"),("5","Inspect","quads / triples / CFG"),("6","Verify","AST vs VM")];
    for i,(n,h,sub) in enumerate(stages):
        x=.65+i*2.08; rect(s,x,2,1.65,2,LIGHT if i not in(3,5) else NAVY,MID if i not in(3,5) else NAVY); box(s,n,x+.12,2.18,.35,.35,14,ORANGE,True); box(s,h,x+.12,2.7,1.4,.3,17,"FFFFFF" if i in(3,5) else NAVY,True,PP_ALIGN.CENTER); box(s,sub,x+.12,3.22,1.4,.3,12,"BFE9E5" if i in(3,5) else MUTED,False,PP_ALIGN.CENTER); box(s,"→",x+1.7,2.75,.35,.35,21,TEAL,True,PP_ALIGN.CENTER) if i<5 else None
    box(s,"Same TAC feeds the human-readable views and the virtual machine.",1,5.25,11.2,.45,20,NAVY,True,PP_ALIGN.CENTER); box(s,"Review-2 extension: CFG analysis sits beside the existing IR formats; it does not replace execution.",1.2,5.95,10.8,.35,13,MUTED,False,PP_ALIGN.CENTER)
    s=p.slides.add_slide(blank); slide_title(s,"Implementation","What is inside the 75% scope",5); foot(s); rows=[("Front end","lexer • parser • AST","DONE"),("IR generation","arithmetic • relational • Boolean • if/while","DONE"),("Optimization","constant fold • copy elimination • cleanup","DONE"),("Verification","AST interpreter • TAC VM • equivalence tests","DONE"),("Review-2 analysis","basic blocks • CFG successor edges","NEW"),("Final 25%","SSA • richer types/features • release pack","NEXT")]; y=1.55
    for a,b,c in rows: rect(s,.85,y,2.25,.65,NAVY if c!="NEXT" else "8A5B20",NAVY); box(s,a,1,y+.16,1.95,.25,14,"FFFFFF",True); rect(s,3.3,y,7.6,.65,LIGHT,MID); box(s,b,3.55,y+.16,6.8,.25,14,INK); box(s,c,11.35,y+.16,1,.25,12,ORANGE if c=="NEXT" else TEAL,True,PP_ALIGN.RIGHT); y+=.78
    box(s,"The implementation boundary is explicit: 75% demonstrated, remaining 25% planned.",1.1,6.55,11.1,.35,16,ORANGE,True,PP_ALIGN.CENTER)
    s=p.slides.add_slide(blank); slide_title(s,"Review-2 addition","From TAC to basic blocks and CFG",6); foot(s); box(s,"leaders",.8,1.45,1.2,.25,12,TEAL,True); box(s,"first instruction • labels • jump targets • instruction after a jump",1.8,1.45,8.5,.25,14,INK); blocks=[("B0","initial assignments","B1"),("B1","while test","B8 / B2"),("B2","Boolean test","B8 / B3"),("B3","if test","B5 / B4"),("B4","else test","B6 / B5"),("B5","then body","B7"),("B6","else body","B7"),("B7","loop back-edge","B1"),("B8","exit + print","exit")]; pos=[(.8,2.15),(4,2.15),(7.2,2.15),(4,3.65),(7.2,3.65),(4,5.15),(7.2,5.15),(1.6,5.15),(10,3.65)]
    for (bid,desc,succ),(x,y) in zip(blocks,pos): rect(s,x,y,2.35,.72,NAVY if bid in("B1","B7") else LIGHT,NAVY); box(s,bid,x+.12,y+.12,.42,.25,13,ORANGE,True); box(s,desc,x+.55,y+.12,1.65,.22,12,"FFFFFF" if bid in("B1","B7") else NAVY,True); box(s,"→ "+succ,x+.55,y+.39,1.65,.18,10,"BFE9E5" if bid in("B1","B7") else MUTED)
    box(s,"Demo result: 9 basic blocks; branch edges, join edges, and the loop back-edge are visible.",.95,6.35,11.2,.35,15,TEAL,True,PP_ALIGN.CENTER)
    s=p.slides.add_slide(blank); slide_title(s,"Live demo","One program, multiple views",7); foot(s); rect(s,.75,1.35,5.6,4.95,"10283E","10283E"); box(s,"$ python -m tacgen examples/07_review2_demo.src --cfg --run",1,1.7,5.1,.3,12,"8FE3DD",False,font="Consolas"); box(s,"L1:\n  if i > n goto L2\n  if sum >= 100 goto L2\n  t1 = i % 2\n  if t1 == 0 goto L3\n  if i != 5 goto L4\nL3: ...\nL2: print sum",1,2.2,5,2.7,15,"FFFFFF",False,font="Consolas"); box(s,"VM output: 43\nVerification: PASS",1,5.35,4.9,.55,18,ORANGE,True,font="Consolas"); rect(s,7,1.35,5.55,4.95,LIGHT,MID); box(s,"What this proves",7.4,1.75,4.6,.3,17,NAVY,True); bullets_ppt(s,["precedence and temporaries","&& / || short-circuit flow","if / else / while labels","CFG derived from the same TAC","AST interpreter agrees with VM"],7.45,2.35,4.5,15,.62)
    s=p.slides.add_slide(blank); slide_title(s,"Evidence","Verification results",8); foot(s); metrics=[("23","tests passing"),("300","random programs"),("4","modes checked"),("43","demo output")]
    for i,(big,small) in enumerate(metrics): x=.8+i*3.05; rect(s,x,1.55,2.5,1.65,NAVY if i in(0,3) else LIGHT,NAVY); box(s,big,x,1.8,2.5,.55,31,"FFFFFF" if i in(0,3) else TEAL,True,PP_ALIGN.CENTER); box(s,small,x+.15,2.5,2.2,.25,13,"BFE9E5" if i in(0,3) else MUTED,False,PP_ALIGN.CENTER)
    box(s,"Cross-check strategy",.85,3.85,2.5,.3,16,TEAL,True); bullets_ppt(s,["run the AST interpreter","run the generated TAC on the VM","compare output and final environment","repeat with optimization and Boolean modes"],.9,4.3,6,16,.52); rect(s,8,3.75,4.3,1.7,"E8F7F5","B7E6E0"); box(s,"AST result = TAC result",8.25,4.15,3.8,.35,20,TEAL,True,PP_ALIGN.CENTER); box(s,"PASS",8.25,4.65,3.8,.35,24,NAVY,True,PP_ALIGN.CENTER)
    s=p.slides.add_slide(blank); slide_title(s,"Boundary","What remains for the final 25%",9); foot(s); bullets_ppt(s,["SSA sketch and phi placement over the CFG","functions, arrays, declarations, and richer type rules","larger benchmarks and optimization study","final screenshots, packaging, and submission polish"],1,1.7,8.8,19,.85); rect(s,9.7,1.65,2.1,2.1,"FFF1DF",ORANGE); box(s,"25%",9.7,2,2.1,.55,30,ORANGE,True,PP_ALIGN.CENTER); box(s,"planned",9.7,2.75,2.1,.3,16,NAVY,True,PP_ALIGN.CENTER); box(s,"Review-2 claim",.95,5.65,2,.3,16,TEAL,True); box(s,"75% demonstrated. Remaining work is visible and intentionally not counted as complete.",3,5.6,8.8,.45,18,NAVY,True)
    s=p.slides.add_slide(blank); slide_title(s,"Q&A","Five answers to keep ready",10); foot(s); qs=[("Why TAC?","Explicit evaluation order and machine-independent IR."),("How does && short-circuit?","False left operand jumps directly to the false label."),("What is a basic block?","A maximal straight-line sequence with one entry."),("How is correctness checked?","AST interpreter and TAC VM must agree."),("Why 75%?","Core pipeline + CFG are done; SSA and final packaging remain.")]
    for i,(q,a) in enumerate(qs): y=1.45+i*.95; box(s,q,.9,y,3.2,.3,16,TEAL,True); box(s,a,4.25,y,7.9,.35,15,INK); rect(s,.9,y+.55,11.3,.01,MID,MID,False)
    box(s,"Thank you",.9,6.35,11.3,.45,25,NAVY,True,PP_ALIGN.CENTER)
    p.save(PPT)

def make_pdf():
    pages=[("THREE ADDRESS CODE GENERATOR","PROJECT REVIEW 2  |  75% IMPLEMENTATION MILESTONE",["BCSE307P Compiler Design Lab","Yash Singh Jhinkwan 24BCI0123  •  Faaz Shabeer 24BCE2754  •  Kevin Gilbert 24BCE2843","Faculty: Dr. Bhawana Tyagi  |  16 October 2026","Assessment IX requires 75% prototype progress, novelty, documentation, and presentation/Q&A."]), ("Executive summary","Implemented and demonstrated",["Lexer, parser/AST, semantic checks, TAC, short-circuit Boolean translation, if/else/while, optimizer, quadruples, triples, and TAC VM.","Review-2 addition: basic-block and control-flow-graph extraction from generated TAC.","Evidence: 23 tests pass; 300 randomized programs × 4 modes pass; demo output is 43 and AST-vs-TAC verification is PASS.","Remaining 25%: SSA sketch, richer language features, larger benchmarks, and final packaging."]), ("Review-2 demonstration","python -m tacgen examples/07_review2_demo.src --cfg --run",["Demo combines while, &&, if/else, ||, modulo, arithmetic temporaries, and print.","CFG result: 9 basic blocks with branch, join, and loop back-edge structure.","VM output: 43","Verification against AST interpreter: PASS","Final variables: n=5, sum=43, i=6"]), ("Verification and next steps","Quality evidence",["23 automated tests passed.","300 randomized programs × 4 optimizer/Boolean-mode combinations passed.","Error paths cover lexical, syntax, semantic, and runtime division-by-zero cases.","The implementation is presented as the requested 75% milestone, not as a finished 100% compiler.","Next: SSA sketch, richer language support, larger benchmarks, and final release documentation."])]
    pdf=fitz.open()
    for h,sub,lines in pages:
        pg=pdf.new_page(width=595,height=842); pg.draw_rect(fitz.Rect(0,0,595,12),color=None,fill=(0,.65,.65)); pg.insert_text((48,65),h,fontsize=24,fontname="hebo",color=(.09,.20,.30)); pg.insert_text((48,93),sub,fontsize=13,fontname="hebo",color=(0,.55,.55)); y=145
        for line in lines:
            pg.draw_circle((58,y-4),3,color=(0,.65,.65),fill=(0,.65,.65)); pg.insert_textbox(fitz.Rect(72,y-18,540,y+34),line,fontsize=12,fontname="helv",color=(.11,.16,.22)); y+=62
        pg.insert_text((48,805),"BCSE307P  •  24BCI0123  •  Project Review 2",fontsize=8,color=(.32,.38,.43))
    pdf.save(PDF)

if __name__ == "__main__":
    make_docx(); make_ppt(); make_pdf(); print(DOCX); print(PDF); print(PPT)
