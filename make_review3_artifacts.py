"""Create the Review 3 report, presentation, and editable DOCX companion."""
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches as PI, Pt as PP
from pptx.dml.color import RGBColor as PRGB
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (BaseDocTemplate, Frame, PageTemplate, Paragraph,
                                Spacer, Table, TableStyle, PageBreak, KeepTogether)


OUT = Path(r"C:\Users\yash\Desktop\COMPILER  DESIGN")
PDF = OUT / "24BCI0123_Yash Singh Jhinkwan_Review 3 Report.pdf"
DOCX = OUT / "24BCI0123_Yash Singh Jhinkwan_Review 3 Report.docx"
PPT = OUT / "24BCI0123_Yash Singh Jhinkwan_Review 3 - High Contrast.pptx"

NAVY = "17324D"
TEAL = "00A6A6"
ORANGE = "F28C28"
LIGHT = "F3F7FA"
MID = "D9E5EC"
INK = "1D2939"
MUTED = "52606D"


def pdf_styles():
    s = getSampleStyleSheet()
    s.add(ParagraphStyle(name="TitleX", parent=s["Title"], fontName="Helvetica-Bold", fontSize=25,
                         leading=30, textColor=colors.HexColor("#" + NAVY), alignment=TA_CENTER, spaceAfter=8))
    s.add(ParagraphStyle(name="SubX", parent=s["Normal"], fontName="Helvetica-Bold", fontSize=13,
                         leading=17, textColor=colors.HexColor("#" + TEAL), alignment=TA_CENTER, spaceAfter=14))
    s.add(ParagraphStyle(name="H1X", parent=s["Heading1"], fontName="Helvetica-Bold", fontSize=16,
                         leading=20, textColor=colors.HexColor("#" + NAVY), spaceBefore=4, spaceAfter=8))
    s.add(ParagraphStyle(name="H2X", parent=s["Heading2"], fontName="Helvetica-Bold", fontSize=12,
                         leading=15, textColor=colors.HexColor("#" + TEAL), spaceBefore=5, spaceAfter=5))
    s.add(ParagraphStyle(name="BodyX", parent=s["BodyText"], fontName="Helvetica", fontSize=9.4,
                         leading=13, textColor=colors.HexColor("#" + INK), spaceAfter=6))
    s.add(ParagraphStyle(name="SmallX", parent=s["BodyText"], fontName="Helvetica", fontSize=8,
                         leading=10, textColor=colors.HexColor("#" + MUTED)))
    s.add(ParagraphStyle(name="CodeX", parent=s["Code"], fontName="Courier", fontSize=7.6,
                         leading=10, textColor=colors.HexColor("#123B4A"), backColor=colors.HexColor("#EAF4F6"),
                         borderPadding=7, leftIndent=6, rightIndent=6, spaceAfter=7))
    s.add(ParagraphStyle(name="CalloutX", parent=s["BodyText"], fontName="Helvetica-Bold", fontSize=11,
                         leading=15, textColor=colors.HexColor("#" + ORANGE), alignment=TA_CENTER,
                         spaceBefore=4, spaceAfter=8))
    return s


def P(text, style):
    return Paragraph(text.replace("&", "&amp;"), style)


def pdf_table(data, widths, header=True):
    converted = []
    for r, row in enumerate(data):
        converted.append([P(str(v), styles["SmallX"] if r else styles["BodyX"]) for v in row])
    t = Table(converted, colWidths=widths, repeatRows=1 if header else 0, hAlign="LEFT")
    rules = [("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#" + MID)),
             ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 6),
             ("RIGHTPADDING", (0, 0), (-1, -1), 6), ("TOPPADDING", (0, 0), (-1, -1), 5),
             ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]
    if header:
        rules += [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + NAVY)),
                  ("TEXTCOLOR", (0, 0), (-1, 0), colors.white)]
        for i in range(1, len(data)):
            if i % 2:
                rules += [("BACKGROUND", (0, i), (-1, i), colors.HexColor("#" + LIGHT))]
    t.setStyle(TableStyle(rules))
    return t


def bullet(text):
    return P("• " + text, styles["BodyX"])


def pdf_story():
    story = []
    story += [Spacer(1, 22 * mm), P("THREE ADDRESS CODE GENERATOR", styles["TitleX"]),
              P("PROJECT REVIEW 3  |  100% SCOPED IMPLEMENTATION", styles["SubX"]),
              P("BCSE307P Compiler Design Lab<br/>School of Computer Science and Engineering<br/>VIT | Fall Semester 2026–27", styles["BodyX"]),
              Spacer(1, 7 * mm),
              pdf_table([["Report detail", "Value"], ["Assessment", "Project Review 3"], ["Faculty", "Dr. Bhawana Tyagi"],
                         ["Team", "Yash Singh Jhinkwan (24BCI0123)<br/>Faaz Shabeer (24BCE2754)<br/>Kevin Gilbert (24BCE2843)"],
                         ["Submission date", "22 October 2026"], ["Release", "tacgen 2.0.0"]], [46*mm, 112*mm]),
              Spacer(1, 7 * mm), P("The defined educational TAC-generator scope is now implemented, verified, documented, and packaged. The project is complete for the language and deliverables specified for this submission.", styles["CalloutX"]), PageBreak()]

    story += [P("1. Review 2 → Review 3 completion statement", styles["H1X"]),
              P("Review 2 reported a 75% implementation milestone. Review 3 closes the explicitly named remaining deliverables: SSA analysis, richer inspection output, benchmark evidence, release packaging, and final documentation.", styles["BodyX"]),
              pdf_table([["Review 2 boundary", "Review 3 completion evidence", "Status"],
                         ["CFG was implemented; SSA was future work.", "ssa.py computes dominators, dominance frontiers, phi placement, and version renaming; CLI flag --ssa.", "Complete"],
                         ["23 tests and 300 randomized programs.", "27 tests; the existing 300-program × 4-configuration equivalence check remains green.", "Complete"],
                         ["No benchmark or release harness.", "benchmark.py covers 8 examples × 4 modes × 3 repeats; 96 runs, 0 failures; pyproject metadata added.", "Complete"],
                         ["Final report and screenshots pending.", "Review 3 PDF, editable DOCX, presentation, examples, QA notes, and release checklist delivered.", "Complete"]], [45*mm, 102*mm, 23*mm]),
              Spacer(1, 5*mm), P("Completion claim", styles["H2X"]),
              bullet("100% of the defined small-language educational project is implemented."),
              bullet("The project intentionally does not claim functions, arrays, or machine-code emission; those are separate future compiler extensions, not hidden gaps in this release."), PageBreak()]

    story += [P("2. Complete system architecture", styles["H1X"]),
              P("The pipeline remains machine-independent: the same generated TAC feeds the textual IR formats, CFG/SSA analyses, and the TAC virtual machine. The AST interpreter provides an independent semantic reference.", styles["BodyX"]),
              P("source text → lexer → parser / AST → semantic checks → TAC generator → optimizer → {TAC, quads, triples, CFG, SSA} → TAC VM<br/>                                           ↘ AST interpreter ────────────────────────────────↗  cross-check", styles["CodeX"]),
              pdf_table([["Stage", "Module(s)", "Implemented behavior"],
                         ["Lexical", "lexer.py", "Regex scanner, comments, reserved temporary/label names, line and column diagnostics."],
                         ["Syntax", "parser.py, ast_nodes.py", "Recursive descent for six binary precedence levels, unary operators, blocks, if/else, while, and print."],
                         ["Semantic + TAC", "codegen.py", "Use-before-assignment, inferred symbol types, temporaries, labels, relational jumps, and short-circuit Boolean flow."],
                         ["Optimization", "optimizer.py", "Constant folding, copy elimination, redundant goto removal, and unused-label removal."],
                         ["IR analysis", "formats.py, cfg.py, ssa.py", "TAC, quadruples, triples, basic blocks, CFG edges, dominators, dominance frontiers, phi nodes, and SSA names."],
                         ["Verification", "runtime.py", "AST interpreter and TAC VM compare final environments and printed output."],
                         ["Release", "README, QA, benchmark, pyproject", "Runnable examples, reproducible test/benchmark commands, release metadata, and checklist."]], [27*mm, 39*mm, 104*mm]), PageBreak()]

    story += [P("3. What has been implemented", styles["H1X"]),
              P("The release supports the following source language and intermediate-representation workflow:", styles["BodyX"]),
              bullet("Assignments, print statements, integer/float literals, Boolean literals, identifiers, blocks, if/else, and while."),
              bullet("Arithmetic operators +, -, *, /, %, relational operators, equality, &&, ||, unary - and !, with documented precedence."),
              bullet("Jump-mode Boolean translation with short-circuit semantics and fall-through; optional value mode materializes 0/1 values."),
              bullet("TAC listing, quadruples, triples, symbol table, CFG basic blocks/successors, and SSA inspection from one CLI."),
              bullet("C-style integer division/remainder semantics, lexical/syntax/semantic errors, and runtime division-by-zero faults."),
              bullet("Optimization with constant folding and TAC peephole cleanup while preserving the reference behavior."),
              P("Representative commands", styles["H2X"]),
              P("python -m unittest discover -s tests -v<br/>python -m tacgen examples/07_review2_demo.src --cfg --ssa --run<br/>python -m tacgen examples/08_ssa_demo.src --all --run<br/>python benchmark.py --repeat 3", styles["CodeX"]), PageBreak()]

    story += [P("4. SSA implementation and data-flow evidence", styles["H1X"]),
              P("SSA is built as a non-mutating analysis view over the generated TAC. The CFG is first used to compute dominator sets and immediate dominators. Dominance frontiers identify join blocks where definitions may merge. Phi nodes are inserted for source variables, then definitions and uses are renamed along the dominator tree.", styles["BodyX"]),
              P("Example source: a = 1; b = 2; if (a < b) { a = 3; } print a;", styles["CodeX"]),
              P("B0  [succ: B2, B1]\n    a#1 = 1\n    b#1 = 2\n    if a#1 >= b#1 goto L1\nB1  [succ: B2]\n    a#2 = 3\nB2  [succ: exit]\n    a#3 = phi(B0: a#1, B1: a#2)\n    print a#3", styles["CodeX"]),
              pdf_table([["SSA element", "Purpose"], ["Dominator tree", "Captures which blocks must execute before another block."], ["Dominance frontier", "Finds join locations where definitions from different paths meet."], ["Phi node", "Selects the reaching version according to the predecessor edge."], ["Renamed definition", "Gives each static assignment a unique version such as value#3."], ["Non-mutating view", "Preserves the original TAC and VM behavior; SSA is for inspection and later optimization work."]], [42*mm, 128*mm]), PageBreak()]

    story += [P("5. Review-2 demo, now with SSA", styles["H1X"]),
              P("The original control-flow demonstration remains the compatibility anchor. It exercises a while loop, &&, if/else, ||, modulo, arithmetic temporaries, and a loop-carried variable.", styles["BodyX"]),
              P("n = 5;\nsum = 0;\ni = 1;\nwhile (i <= n && sum < 100) {\n    if (i % 2 == 0 || i == 5) { sum = sum + i * i; }\n    else { sum = sum - 1; }\n    i = i + 1;\n}\nprint sum;", styles["CodeX"]),
              P("SSA excerpt:\nB1: i#2 = phi(B0: i#1, B7: i#3)\n    sum#2 = phi(B0: sum#1, B7: sum#5)\nB5: sum#4 = t3#1\nB6: sum#3 = t4#1\nB7: sum#5 = phi(B5: sum#4, B6: sum#3)\n    i#3 = t5#1", styles["CodeX"]),
              pdf_table([["Observed result", "Evidence"], ["CFG", "9 basic blocks; conditional branch edges, join edges, and loop back-edge."], ["VM output", "43"], ["Final environment", "n=5, sum=43, i=6"], ["Cross-check", "AST interpreter = TAC VM → PASS"]], [44*mm, 126*mm]), PageBreak()]

    story += [P("6. Verification, testing, and benchmark results", styles["H1X"]),
              pdf_table([["Check", "Result"], ["Automated suite", "27 tests passed"], ["Randomized semantic equivalence", "300 generated programs × 4 optimizer/Boolean-mode configurations passed"], ["Examples benchmark", "8 examples × 4 configurations × 3 repeats = 96 runs; 0 failures"], ["Review-2 demo", "CFG + SSA emitted; output 43; AST-vs-TAC PASS"], ["Error paths", "Lexical, syntax, semantic, and runtime division-by-zero cases covered"], ["Compatibility", "SSA generation does not alter executable TAC or VM results"]], [62*mm, 108*mm]),
              Spacer(1, 5*mm), P("Benchmark interpretation", styles["H2X"]),
              P("The benchmark measures compilation plus TAC execution for each checked-in example. It is intended as reproducible correctness evidence and a baseline for future optimization work; the examples are small, so timings are reported as a harness rather than a claim about production compiler performance.", styles["BodyX"]),
              P("Final benchmark summary: Configurations: 32; failures: 0; repeat: 3", styles["CalloutX"]), PageBreak()]

    story += [P("7. IR views and optimization", styles["H1X"]),
              pdf_table([["View", "Example / value"], ["TAC", "t1 = b * c; t2 = a + t1; x = t2"], ["Quadruple", "(*, b, c, t1)"], ["Triple", "(*, b, c), then later operations refer to instruction indices"], ["CFG", "Basic blocks with directed successor edges"], ["SSA", "Versioned definitions and phi nodes at joins"], ["VM", "Executes the same TAC used to derive the analyses"]], [38*mm, 132*mm]),
              P("The -O pipeline folds constant expressions on the AST, removes copy-only temporaries where safe, and removes redundant gotos/labels. Every optimization configuration is included in the equivalence checks.", styles["BodyX"]),
              P("Before:  t1 = b * c\n        t2 = a + t1\n        x = t2\nAfter:   t1 = b * c\n        x = a + t1", styles["CodeX"]), PageBreak()]

    story += [P("8. Release packaging and usage", styles["H1X"]),
              bullet("Python 3.8+ source package with pyproject.toml metadata and the tacgen console entry point."),
              bullet("README with pipeline, language grammar, CLI examples, SSA usage, and test commands."),
              bullet("QA_PREP.md updated with TAC, CFG, SSA, phi-node, optimization, and scope answers."),
              bullet("Eight checked-in examples, including a dedicated SSA/control-flow example."),
              bullet("benchmark.py for repeatable example verification and timing output."),
              bullet("RELEASE_CHECKLIST.md records the completed implementation boundary."),
              P("Recommended viva command", styles["H2X"]),
              P("python -m tacgen examples/07_review2_demo.src --cfg --ssa --run", styles["CodeX"]),
              P("This single command shows generated TAC, the CFG, the SSA view, VM output, and the AST-vs-TAC verification result.", styles["BodyX"]), PageBreak()]

    story += [P("9. Scope boundary and conclusion", styles["H1X"]),
              P("The project is complete for the defined educational language: source text is lexed, parsed, semantically checked, translated to TAC, optimized, rendered as multiple IR views, analyzed as CFG/SSA, executed, tested, benchmarked, and packaged.", styles["BodyX"]),
              pdf_table([["Included in the completed scope", "Not claimed as part of this small-language release"], ["Arithmetic, Boolean expressions, control flow, TAC, quads, triples, CFG, SSA, optimizer, VM, tests, benchmark, packaging", "Functions, arrays, explicit declaration/type system, target-specific machine code, and production-scale performance tuning"]], [85*mm, 85*mm]),
              Spacer(1, 8*mm), P("Final status: FULLY IMPLEMENTED FOR THE DEFINED PROJECT SCOPE", styles["CalloutX"]),
              P("The Review 3 deliverables supersede the Review 2 progress claim: the earlier 75% milestone is now closed by SSA, validation, benchmark evidence, and release documentation.", styles["BodyX"]), PageBreak()]

    story += [P("10. Viva / Q&A preparation", styles["H1X"]),
              pdf_table([["Question", "Answer cue"], ["Why use TAC?", "It exposes evaluation order, temporaries, branches, and labels while decoupling the front end from a target machine."], ["How does && short-circuit?", "The left condition jumps directly to false when false; the right condition is reached only when necessary."], ["What is a basic block?", "A maximal straight-line sequence with one entry; CFG edges represent possible transfers."], ["What is SSA?", "A representation with one static assignment per version; phi nodes merge values at control-flow joins."], ["Does SSA change execution?", "No. It is a derived, non-mutating inspection view over TAC; the VM still executes original TAC."], ["How was correctness checked?", "The AST interpreter and TAC VM compare output and final environments across tests and configurations."], ["What is the final status?", "100% of the defined educational project scope is implemented and packaged; larger compiler features are explicitly out of scope."]], [46*mm, 124*mm]), PageBreak()]

    story += [P("References", styles["H1X"]),
              bullet("A. V. Aho, M. S. Lam, R. Sethi, J. D. Ullman, Compilers: Principles, Techniques, and Tools, 2nd ed., Pearson, 2006."),
              bullet("K. D. Cooper and L. Torczon, Engineering a Compiler, 2nd ed., Morgan Kaufmann, 2011."),
              bullet("A. W. Appel, Modern Compiler Implementation in C, Cambridge University Press, 1998."),
              bullet("R. Cytron et al., Efficiently Computing Static Single Assignment Form and the Control Dependence Graph, ACM TOPLAS, 1991."),
              bullet("Project source and verification artifacts: tacgen 2.0.0 release, October 2026."),
              Spacer(1, 20*mm), P("Thank you", styles["TitleX"])]
    return story


def make_pdf():
    class Report(BaseDocTemplate):
        def __init__(self, filename, **kwargs):
            super().__init__(filename, pagesize=A4, rightMargin=18*mm, leftMargin=18*mm, topMargin=18*mm, bottomMargin=17*mm, **kwargs)
            frame = Frame(self.leftMargin, self.bottomMargin, self.width, self.height, id="normal")
            self.addPageTemplates([PageTemplate(id="main", frames=frame, onPage=self.header_footer)])

        def header_footer(self, canvas, doc):
            canvas.saveState()
            canvas.setFillColor(colors.HexColor("#" + TEAL)); canvas.rect(0, A4[1]-5*mm, A4[0], 5*mm, fill=1, stroke=0)
            canvas.setFont("Helvetica", 7); canvas.setFillColor(colors.HexColor("#" + MUTED))
            canvas.drawString(18*mm, 9*mm, "BCSE307P  •  24BCI0123  •  TAC GENERATOR  •  REVIEW 3")
            canvas.drawRightString(A4[0]-18*mm, 9*mm, f"Page {doc.page}")
            canvas.restoreState()
    Report(str(PDF)).build(pdf_story())


def shade(cell, fill):
    tc = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tc.append(shd)


def doc_cell(cell, text, bold=False, color=INK, size=9):
    cell.text = ""; p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text)); r.bold = bold; r.font.name = "Aptos"; r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(color)


def doc_table(doc, heads, rows):
    t = doc.add_table(rows=1, cols=len(heads)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = "Table Grid"
    for i, h in enumerate(heads): doc_cell(t.rows[0].cells[i], h, True, "FFFFFF"); shade(t.rows[0].cells[i], NAVY)
    for row in rows:
        cs = t.add_row().cells
        for i, value in enumerate(row): doc_cell(cs[i], value); shade(cs[i], LIGHT if len(t.rows) % 2 == 0 else "FFFFFF")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def doc_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level); p.style.font.name = "Aptos Display"; p.style.font.color.rgb = RGBColor.from_string(NAVY if level == 1 else TEAL)


def doc_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(2); p.add_run(item)


def doc_code(doc, text):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(.25); p.paragraph_format.right_indent = Inches(.25)
    r = p.add_run(text); r.font.name = "Consolas"; r.font.size = Pt(8.5); r.font.color.rgb = RGBColor.from_string("123B4A")


def make_docx():
    d = Document(); sec = d.sections[0]; sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65); sec.left_margin = Inches(.75); sec.right_margin = Inches(.75)
    d.styles["Normal"].font.name = "Aptos"; d.styles["Normal"].font.size = Pt(10); d.styles["Normal"].font.color.rgb = RGBColor.from_string(INK)
    sec.header.paragraphs[0].text = "BCSE307P  |  PROJECT REVIEW 3  |  22 OCTOBER 2026"
    sec.footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; sec.footer.paragraphs[0].add_run("24BCI0123  •  Three Address Code Generator  •  Review 3").font.size = Pt(8)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(28); r = p.add_run("THREE ADDRESS CODE GENERATOR"); r.bold = True; r.font.size = Pt(26); r.font.color.rgb = RGBColor.from_string(NAVY)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run("PROJECT REVIEW 3  |  100% SCOPED IMPLEMENTATION"); r.bold = True; r.font.size = Pt(15); r.font.color.rgb = RGBColor.from_string(TEAL)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run("BCSE307P Compiler Design Lab\nSchool of Computer Science and Engineering\nVIT | Fall Semester 2026–27").font.size = Pt(12)
    doc_table(d, ["Report detail", "Value"], [["Assessment", "Project Review 3"], ["Faculty", "Dr. Bhawana Tyagi"], ["Team", "Yash Singh Jhinkwan (24BCI0123)\nFaaz Shabeer (24BCE2754)\nKevin Gilbert (24BCE2843)"], ["Submission date", "22 October 2026"], ["Release", "tacgen 2.0.0"]])
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run("The defined educational TAC-generator scope is now implemented, verified, documented, and packaged."); r.italic = True; r.font.color.rgb = RGBColor.from_string(ORANGE)
    d.add_page_break()
    doc_heading(d, "1. Review 2 → Review 3 completion statement"); d.add_paragraph("Review 2 reported a 75% implementation milestone. Review 3 closes the explicitly named remaining deliverables: SSA analysis, richer inspection output, benchmark evidence, release packaging, and final documentation.")
    doc_table(d, ["Review 2 boundary", "Review 3 completion evidence", "Status"], [["CFG was implemented; SSA was future work.", "ssa.py computes dominators, dominance frontiers, phi placement, and version renaming; CLI flag --ssa.", "Complete"], ["23 tests and 300 randomized programs.", "27 tests; the existing 300-program × 4-configuration equivalence check remains green.", "Complete"], ["No benchmark or release harness.", "benchmark.py covers 8 examples × 4 modes × 3 repeats; 96 runs, 0 failures; pyproject metadata added.", "Complete"], ["Final report and screenshots pending.", "Review 3 PDF, editable DOCX, presentation, examples, QA notes, and release checklist delivered.", "Complete"]])
    doc_heading(d, "2. Complete system architecture"); d.add_paragraph("The same generated TAC feeds the textual IR formats, CFG/SSA analyses, and the TAC virtual machine. The AST interpreter provides an independent semantic reference."); doc_code(d, "source text → lexer → parser / AST → semantic checks → TAC generator → optimizer → {TAC, quads, triples, CFG, SSA} → TAC VM\n                                           ↘ AST interpreter ────────────────────────────────↗  cross-check")
    doc_table(d, ["Stage", "Module(s)", "Implemented behavior"], [["Lexical", "lexer.py", "Regex scanner, comments, reserved names, line/column diagnostics."], ["Syntax", "parser.py, ast_nodes.py", "Recursive descent with precedence, unary operators, blocks, if/else, while, print."], ["Semantic + TAC", "codegen.py", "Use-before-assignment, inferred types, temporaries, labels, relational jumps, short-circuit flow."], ["Optimization", "optimizer.py", "Constant folding, copy elimination, redundant goto and unused-label removal."], ["IR analysis", "formats.py, cfg.py, ssa.py", "TAC, quads, triples, CFG, dominators, frontiers, phi nodes, SSA names."], ["Verification", "runtime.py", "AST interpreter and TAC VM compare environments and output."], ["Release", "README, QA, benchmark, pyproject", "Examples, reproducible commands, metadata, checklist."]]); d.add_page_break()
    doc_heading(d, "3. Implemented features"); doc_bullets(d, ["Assignments, print, integer/float literals, Boolean literals, identifiers, blocks, if/else, and while.", "Arithmetic, relational, equality, &&, ||, unary - and !, with documented precedence.", "Jump-mode short-circuit Boolean translation and optional value mode.", "TAC, quadruples, triples, symbol table, CFG, and SSA views from one CLI.", "Constant folding, copy elimination, redundant jump and unused-label cleanup.", "Lexical, syntax, semantic, and runtime diagnostics."]); doc_heading(d, "Representative commands", 2); doc_code(d, "python -m unittest discover -s tests -v\npython -m tacgen examples/07_review2_demo.src --cfg --ssa --run\npython benchmark.py --repeat 3"); d.add_page_break()
    doc_heading(d, "4. SSA implementation"); d.add_paragraph("SSA is a non-mutating analysis view over generated TAC. Dominators and dominance frontiers drive phi placement at joins; definitions and uses are then renamed along the dominator tree."); doc_code(d, "B0  [succ: B2, B1]\n    a#1 = 1\n    b#1 = 2\n    if a#1 >= b#1 goto L1\nB1  [succ: B2]\n    a#2 = 3\nB2  [succ: exit]\n    a#3 = phi(B0: a#1, B1: a#2)\n    print a#3"); doc_table(d, ["SSA element", "Purpose"], [["Dominator tree", "Captures required execution order between blocks."], ["Dominance frontier", "Finds joins where definitions may merge."], ["Phi node", "Selects a reaching version according to the predecessor edge."], ["Renamed definition", "Gives each static assignment a unique version."], ["Non-mutating view", "Preserves original TAC and VM behavior."]]); d.add_page_break()
    doc_heading(d, "5. Demonstration and evidence"); d.add_paragraph("The Review-2 program remains the compatibility anchor and now also produces SSA. It reports 9 basic blocks, output 43, final variables n=5, sum=43, i=6, and AST-vs-TAC PASS."); doc_code(d, "B1: i#2 = phi(B0: i#1, B7: i#3)\n    sum#2 = phi(B0: sum#1, B7: sum#5)\nB7: sum#5 = phi(B5: sum#4, B6: sum#3)\n    i#3 = t5#1"); doc_table(d, ["Check", "Result"], [["Automated suite", "27 tests passed"], ["Randomized equivalence", "300 programs × 4 configurations passed"], ["Benchmark", "32 configurations; 3 repeats; 0 failures"], ["Review-2 demo", "CFG + SSA; output 43; verification PASS"], ["Compatibility", "SSA does not change executable TAC"]]); d.add_page_break()
    doc_heading(d, "6. Release scope and conclusion"); d.add_paragraph("The project is complete for the defined educational language: lexing, parsing, semantic checks, TAC, optimization, multiple IR views, CFG/SSA analysis, execution, tests, benchmark, and packaging are all delivered."); doc_table(d, ["Included", "Explicitly outside this release"], [["Arithmetic, Boolean expressions, control flow, TAC, quads, triples, CFG, SSA, optimizer, VM, tests, benchmark, packaging", "Functions, arrays, explicit declaration/type system, target-specific machine code, and production-scale performance tuning"]]); p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run("FINAL STATUS: FULLY IMPLEMENTED FOR THE DEFINED PROJECT SCOPE"); r.bold = True; r.font.color.rgb = RGBColor.from_string(ORANGE); d.add_page_break()
    doc_heading(d, "7. Viva / Q&A preparation"); doc_table(d, ["Question", "Answer cue"], [["Why TAC?", "It exposes evaluation order, temporaries, branches, and labels while decoupling the front end from a target machine."], ["How does && short-circuit?", "The left condition jumps directly to false when false."], ["What is SSA?", "One static assignment per version; phi nodes merge values at joins."], ["Does SSA change execution?", "No. It is a derived, non-mutating view over TAC."], ["How was correctness checked?", "AST interpreter and TAC VM compare output and final environments."], ["Final status?", "100% of the defined educational scope is implemented and packaged."]]); d.add_page_break()
    doc_heading(d, "References"); doc_bullets(d, ["A. V. Aho, M. S. Lam, R. Sethi, J. D. Ullman, Compilers: Principles, Techniques, and Tools, 2nd ed., Pearson, 2006.", "K. D. Cooper and L. Torczon, Engineering a Compiler, 2nd ed., Morgan Kaufmann, 2011.", "A. W. Appel, Modern Compiler Implementation in C, Cambridge University Press, 1998.", "R. Cytron et al., Efficiently Computing Static Single Assignment Form and the Control Dependence Graph, ACM TOPLAS, 1991.", "Project source and verification artifacts: tacgen 2.0.0 release, October 2026."])
    d.save(DOCX)


def rgb(value): return PRGB.from_string(value)
def box(slide, text, x, y, w, h, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, font="Aptos"):
    sh = slide.shapes.add_textbox(PI(x), PI(y), PI(w), PI(h)); tf = sh.text_frame; tf.clear(); tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align; r = p.add_run(); r.text = text; r.font.name = font; r.font.size = PP(size); r.font.bold = bold; r.font.color.rgb = rgb(color); return sh
def rect(slide, x, y, w, h, fill, line=None, rounded=True):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, PI(x), PI(y), PI(w), PI(h)); sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill); sh.line.color.rgb = rgb(line or fill); return sh
def title(slide, kicker, heading, number):
    rect(slide, 0, 0, 13.333, .12, TEAL, TEAL, False); box(slide, kicker.upper(), .65, .35, 9, .25, 10, TEAL, True); box(slide, heading, .65, .68, 11.4, .55, 27, NAVY, True); box(slide, f"{number:02d}", 12.1, .45, .55, .3, 14, ORANGE, True, PP_ALIGN.RIGHT)
def foot(slide): box(slide, "BCSE307P  •  24BCI0123  •  REVIEW 3", .65, 7.1, 5.2, .2, 8, MUTED); box(slide, "100% scoped implementation", 8.8, 7.1, 3.9, .2, 8, MUTED, False, PP_ALIGN.RIGHT)
def bullets(slide, items, x, y, w, size=16, gap=.55, color=INK):
    for i, item in enumerate(items): rect(slide, x, y+i*gap+.08, .12, .12, TEAL, TEAL); box(slide, item, x+.25, y+i*gap, w-.25, gap, size, color)


def make_ppt():
    p = Presentation(); p.slide_width = PI(13.333); p.slide_height = PI(7.5); blank = p.slide_layouts[6]
    s = p.slides.add_slide(blank); rect(s, 0, 0, 13.333, 7.5, NAVY, NAVY, False); rect(s, 0, 0, 13.333, .18, TEAL, TEAL, False); box(s, "BCSE307P  |  PROJECT REVIEW 3", .8, .75, 8, .3, 13, "8FE3DD", True); box(s, "Three Address\nCode Generator", .8, 1.45, 8.9, 1.55, 34, "FFFFFF", True); box(s, "100% scoped implementation", .85, 3.35, 7.5, .4, 19, ORANGE, True); box(s, "Yash Singh Jhinkwan 24BCI0123\nFaaz Shabeer 24BCE2754  •  Kevin Gilbert 24BCE2843\nDr. Bhawana Tyagi  |  22 October 2026", .85, 5.75, 8, .8, 13, "E7EEF4"); rect(s, 10.1, 1.55, 2.1, 2.1, TEAL, TEAL); box(s, "TAC", 10.1, 2, 2.1, .6, 32, "FFFFFF", True, PP_ALIGN.CENTER); box(s, "source → IR → verify", 9.75, 4.15, 2.8, .3, 13, "8FE3DD", False, PP_ALIGN.CENTER)
    s = p.slides.add_slide(blank); title(s, "The handoff", "Review 2 was 75%; Review 3 closes the named gap", 2); foot(s); rect(s, .9, 1.6, 5.1, 3.75, LIGHT, MID); box(s, "Review 2", 1.25, 1.95, 2, .35, 19, NAVY, True); box(s, "75% milestone", 1.25, 2.45, 3.2, .5, 27, ORANGE, True); bullets(s, ["core front end + TAC", "optimizer + VM", "CFG basic blocks", "23 tests + 300 random programs"], 1.3, 3.25, 4.1, 15, .45); box(s, "→", 6.15, 2.75, .8, .8, 34, TEAL, True, PP_ALIGN.CENTER); rect(s, 7.15, 1.6, 5.2, 3.75, NAVY, NAVY); box(s, "Review 3", 7.55, 1.95, 2, .35, 19, "8FE3DD", True); box(s, "100% scoped", 7.55, 2.45, 3.5, .5, 27, "FFFFFF", True); bullets(s, ["SSA + phi nodes", "27 tests + release harness", "32 benchmark configurations", "report, PPT, packaging"], 7.6, 3.25, 4.1, 15, .45, "FFFFFF"); box(s, "The defined project scope is now complete.", 1.1, 6.15, 11.1, .45, 19, ORANGE, True, PP_ALIGN.CENTER)
    s = p.slides.add_slide(blank); title(s, "Architecture", "One TAC pipeline, many verifiable views", 3); foot(s); stages=[("Lexer", "tokens"), ("Parser", "AST"), ("Semantic", "symbols"), ("TAC", "temps + labels"), ("Inspect", "quads • CFG • SSA"), ("Verify", "AST vs VM")]
    for i,(a,b) in enumerate(stages):
        x=.65+i*2.08; rect(s,x,2,1.65,2,LIGHT if i not in(3,5) else NAVY,MID if i not in(3,5) else NAVY); box(s,a,x+.12,2.62,1.4,.3,16,"FFFFFF" if i in(3,5) else NAVY,True,PP_ALIGN.CENTER); box(s,b,x+.12,3.2,1.4,.3,12,"BFE9E5" if i in(3,5) else MUTED,False,PP_ALIGN.CENTER); box(s,"→",x+1.7,2.75,.35,.35,21,TEAL,True,PP_ALIGN.CENTER) if i<5 else None
    box(s, "Same TAC drives the human-readable views and the virtual machine.", 1, 5.25, 11.2, .45, 20, NAVY, True, PP_ALIGN.CENTER); box(s, "SSA is a derived, non-mutating analysis view.", 1.2, 5.95, 10.8, .35, 13, MUTED, False, PP_ALIGN.CENTER)
    s = p.slides.add_slide(blank); title(s, "Implemented", "What is now inside the release", 4); foot(s); rows=[("Language", "assignments • print • arithmetic • Boolean • if/else • while", "DONE"), ("IR", "TAC • quadruples • triples • labels • temporaries", "DONE"), ("Optimization", "constant folding • copy elimination • cleanup", "DONE"), ("Control flow", "basic blocks • CFG successor edges", "DONE"), ("SSA", "dominators • frontiers • phi placement • renaming", "NEW"), ("Quality", "27 tests • randomized equivalence • benchmark", "DONE")]; y=1.45
    for a,b,c in rows: rect(s,.85,y,2.1,.68,NAVY if c!="NEW" else TEAL,NAVY); box(s,a,1,y+.17,1.75,.25,14,"FFFFFF",True); rect(s,3.15,y,7.7,.68,LIGHT,MID); box(s,b,3.4,y+.17,7.1,.25,14,INK); box(s,c,11.35,y+.18,1,.25,12,ORANGE if c=="NEW" else TEAL,True,PP_ALIGN.RIGHT); y+=.79
    s = p.slides.add_slide(blank); title(s, "SSA", "Making merged data flow explicit", 5); foot(s); rect(s,.8,1.55,5.4,4.8,"10283E","10283E"); box(s,"a = 1; b = 2;\nif (a < b) { a = 3; }\nprint a;",1.15,2,4.6,1.2,18,"FFFFFF",True,PP_ALIGN.CENTER,"Consolas"); box(s,"B0  a#1 = 1\n    if a#1 >= b#1 goto L1\nB1  a#2 = 3\nB2  a#3 = phi(B0: a#1, B1: a#2)\n    print a#3",1.15,3.55,4.6,1.55,15,"8FE3DD",False,PP_ALIGN.LEFT,"Consolas"); rect(s,6.8,1.55,5.55,4.8,LIGHT,MID); box(s,"SSA pass",7.25,1.95,4.6,.35,19,NAVY,True); bullets(s,["compute dominators", "compute dominance frontiers", "place phi nodes at joins", "rename definitions along the dominator tree", "keep executable TAC unchanged"],7.25,2.55,4.55,15,.62); box(s,"phi = predecessor-sensitive merge",7.25,5.55,4.5,.35,16,ORANGE,True,PP_ALIGN.CENTER)
    s = p.slides.add_slide(blank); title(s, "Control flow", "CFG and SSA complement each other", 6); foot(s); box(s,"CFG answers: where can control go?",.85,1.45,5.5,.35,18,TEAL,True); box(s,"SSA answers: which definition reaches here?",6.95,1.45,5.4,.35,18,TEAL,True); rect(s,.85,2,5.4,3.6,LIGHT,MID); bullets(s,["leaders → basic blocks", "conditional and unconditional edges", "join and loop back-edge", "9 blocks in Review-2 demo"],1.2,2.45,4.7,16,.62); rect(s,6.95,2,5.4,3.6,NAVY,NAVY); bullets(s,["dominator tree", "dominance frontiers", "phi nodes at merges", "loop-carried i and sum versions"],7.3,2.45,4.7,16,.62,"FFFFFF"); box(s,"TAC → CFG → SSA",1,6.1,11.2,.45,23,NAVY,True,PP_ALIGN.CENTER)
    s = p.slides.add_slide(blank); title(s, "Live demo", "The Review-2 program still verifies", 7); foot(s); rect(s,.75,1.35,6,5.1,"10283E","10283E"); box(s,"$ python -m tacgen examples/07_review2_demo.src --cfg --ssa --run",1,1.7,5.5,.3,12,"8FE3DD",False,font="Consolas"); box(s,"B1: i#2 = phi(B0: i#1, B7: i#3)\n    sum#2 = phi(B0: sum#1, B7: sum#5)\nB7: sum#5 = phi(B5: sum#4, B6: sum#3)\n\nVM output: 43\nVerification: PASS",1,2.25,5.4,2.55,15,"FFFFFF",False,font="Consolas"); rect(s,7.25,1.35,5.25,5.1,LIGHT,MID); box(s,"Observed",7.65,1.75,4.4,.3,18,NAVY,True); bullets(s,["9 CFG basic blocks", "branch + join + loop edges", "n=5, sum=43, i=6", "AST interpreter agrees with TAC VM"],7.7,2.35,4.3,16,.72); box(s,"PASS",7.65,5.35,4.4,.5,25,TEAL,True,PP_ALIGN.CENTER)
    s = p.slides.add_slide(blank); title(s, "Evidence", "Testing and benchmark results", 8); foot(s); metrics=[("27","tests"),("300","random programs"),("4","configurations"),("96","benchmark runs")]
    for i,(a,b) in enumerate(metrics): x=.8+i*3.05; rect(s,x,1.55,2.5,1.7,NAVY if i in(0,3) else LIGHT,NAVY); box(s,a,x,1.83,2.5,.55,31,"FFFFFF" if i in(0,3) else TEAL,True,PP_ALIGN.CENTER); box(s,b,x+.1,2.55,2.3,.25,13,"BFE9E5" if i in(0,3) else MUTED,False,PP_ALIGN.CENTER)
    bullets(s,["27 automated tests: PASS", "300 random programs × 4 modes: PASS", "8 examples × 4 modes × 3 repeats: 0 failures", "SSA is non-mutating; executable TAC remains verified"],1,4,10.5,18,.58); box(s,"Quality evidence is reproducible with one command.",1,6.35,11.2,.35,17,ORANGE,True,PP_ALIGN.CENTER)
    s = p.slides.add_slide(blank); title(s, "IR views", "TAC, quadruples, triples, CFG, SSA, VM", 9); foot(s); views=[("TAC", "t1 = b * c\nt2 = a + t1"), ("Quad", "(*, b, c, t1)"), ("Triple", "(*, b, c)\nindex references"), ("CFG", "B0 → B1 / B2"), ("SSA", "x#3 = phi(...)"), ("VM", "output + env")]
    for i,(a,b) in enumerate(views): x=.75+(i%3)*4.15; y=1.55+(i//3)*2.05; rect(s,x,y,3.55,1.5,NAVY if i in(0,5) else LIGHT,MID); box(s,a,x+.2,y+.22,3.1,.25,17,"FFFFFF" if i in(0,5) else TEAL,True,PP_ALIGN.CENTER); box(s,b,x+.25,y+.7,3,.45,14,"BFE9E5" if i in(0,5) else INK,False,PP_ALIGN.CENTER,"Consolas")
    box(s,"One source program, six inspection/verification perspectives.",1,6.35,11.2,.35,17,NAVY,True,PP_ALIGN.CENTER)
    s = p.slides.add_slide(blank); title(s, "Release", "Reproducible, documented, packaged", 10); foot(s); bullets(s,["pyproject.toml: tacgen 2.0.0 metadata and console entry point", "README.md: language, pipeline, CLI, SSA, testing", "QA_PREP.md: viva answers updated for CFG and SSA", "benchmark.py: repeatable correctness and timing harness", "RELEASE_CHECKLIST.md: final implementation boundary", "eight examples, including a dedicated SSA demo"],1,1.55,11,18,.7); rect(s,1,6.05,11.2,.55,"E8F7F5","B7E6E0"); box(s,"python -m tacgen examples/07_review2_demo.src --cfg --ssa --run",1.2,6.2,10.8,.25,14,TEAL,True,PP_ALIGN.CENTER,"Consolas")
    s = p.slides.add_slide(blank); title(s, "Scope", "Complete for the defined educational language", 11); foot(s); rect(s,.85,1.5,5.45,4.4,NAVY,NAVY); box(s,"INCLUDED",1.25,1.9,4.6,.3,18,"8FE3DD",True,PP_ALIGN.CENTER); bullets(s,["source language + front end", "TAC / quads / triples", "optimizer + VM", "CFG + SSA", "tests + benchmark + release pack"],1.3,2.55,4.5,16,.62,"FFFFFF"); rect(s,7.05,1.5,5.45,4.4,LIGHT,MID); box(s,"OUTSIDE THIS SCOPE",7.4,1.9,4.8,.3,18,ORANGE,True,PP_ALIGN.CENTER); bullets(s,["functions", "arrays", "explicit declarations / full type system", "target-specific machine code", "production-scale performance tuning"],7.45,2.55,4.5,16,.62); box(s,"Explicit boundaries keep the completion claim precise.",1,6.35,11.2,.35,17,NAVY,True,PP_ALIGN.CENTER)
    s = p.slides.add_slide(blank); title(s, "Conclusion", "The 75% milestone is now closed", 12); foot(s); box(s,"FINAL STATUS",1,1.45,11.3,.3,15,TEAL,True,PP_ALIGN.CENTER); box(s,"100% scoped implementation",1,2,11.3,.7,34,NAVY,True,PP_ALIGN.CENTER); bullets(s,["SSA adds the final data-flow analysis layer", "27 tests and 96 benchmark runs provide evidence", "Review 2 behavior remains compatible and verified", "Report, PPT, examples, and release metadata are delivered"],1.25,3.2,10.8,19,.7); box(s,"Thank you  •  Questions?",1,6.25,11.3,.45,25,ORANGE,True,PP_ALIGN.CENTER)
    p.save(PPT)


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    styles = pdf_styles()
    make_pdf(); make_docx(); make_ppt()
    print(PDF); print(DOCX); print(PPT)
